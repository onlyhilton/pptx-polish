#!/usr/bin/env python3
"""normalize_deck.py — Step 1 规范化：标题归位到占位符 + 字体继承主题。

解决的具体问题：原稿的标题常用自由文本框写就，不继承版式。换模板时文本框不跟随
版式变化，导致大面积错位（如黑底变白底）。本脚本把标题搬进 title 占位符，并清除
run 上的硬编码字体，让文字继承主题字体。

依赖：python-pptx + lxml

用法：
    python normalize_deck.py "<in.pptx>" --out "<out.pptx>" [--dry-run]

产物：
    <out.pptx>            规范化后的新文件（原件不动）
    <out>_normalize.json  变更清单
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

# 标题判定规则与审计脚本共用同一份实现，避免"改了一边、漏了另一边"
from deck_rules import (  # noqa: E402
    COVER_ZONE_PX, MIN_TITLE_PT, TITLE_MAX_CHARS, TITLE_ZONE_PX,
    emu_to_px, max_font_pt, title_zone_for,
)
from deck_rules import find_title_textbox as _shared_find_title

PX = 9525
TITLE_PH_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def emu(px_value: float) -> int:
    return int(px_value * PX)


def iter_shapes(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in iter_shapes(shp.shapes):
                yield sub


def iter_text_frames(shape):
    """产出形状内的所有 text_frame（含表格单元格）。"""
    try:
        if shape.has_text_frame:
            yield shape.text_frame
    except Exception:
        pass
    try:
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
    except Exception:
        pass


def find_title_placeholder(slide):
    for shp in iter_shapes(slide.shapes):
        try:
            if shp.is_placeholder and shp.placeholder_format.type in TITLE_PH_TYPES:
                return shp
        except Exception:
            continue
    for shp in iter_shapes(slide.shapes):
        try:
            if shp.is_placeholder and shp.placeholder_format.idx == 0:
                return shp
        except Exception:
            continue
    return None


def find_title_textbox(slide, has_title_ph_with_text: bool, zone: str = "top"):
    """找疑似标题的自由文本框 —— 判定规则见 deck_rules（与审计脚本共用）。

    zone='top'  内容版式：只在上部区域找，字号无从判断时也认紧贴顶部的短文本
    zone='full' 封面版式：全页找，取字号最大的短文本
    zone=None   该版式无标题槽位：不归位
    """
    return _shared_find_title(iter_shapes(slide.shapes), zone,
                              has_title_ph_with_text)


TITLE_SP_XML = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="{sid}" name="Title Placeholder {sid}"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="{phtype}"/></p:nvPr></p:nvSpPr>'
    '<p:spPr/>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>'
    '</p:sp>'
)


def add_title_placeholder(slide, ph_type: str = "title", next_id: int = 900):
    """新建标题占位符。

    刻意不写 <a:xfrm>：位置与尺寸完全继承版式，这样以后再换模板，
    标题会自动跟着新版式走 —— 这正是"规范化"要买到的保险。
    """
    from pptx.oxml import parse_xml
    used = {sp.get("id") for sp in slide.shapes._spTree.iter(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")}
    while str(next_id) in used:
        next_id += 1
    sp = parse_xml(TITLE_SP_XML.format(sid=next_id, phtype=ph_type))
    spTree = slide.shapes._spTree
    spTree.insert(2, sp)                       # nvGrpSpPr / grpSpPr 之后
    return find_title_placeholder(slide)


def clear_hardcoded_fonts(root_element, counter: list):
    """清除 run / 段落默认属性上的字体指定，使文字继承主题。"""
    from pptx.oxml.ns import qn
    for rPr in root_element.iter(qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")):
        for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
            el = rPr.find(qn(tag))
            if el is not None:
                rPr.remove(el)
                counter[0] += 1


def title_ph_type_for(slide) -> str:
    """新建占位符时该用什么类型：版式有 ctrTitle 就用 ctrTitle，否则 title。"""
    try:
        for ph in slide.slide_layout.placeholders:
            try:
                if ph.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE:
                    return "ctrTitle"
            except Exception:
                continue
    except Exception:
        pass
    return "title"


def process_item(slide, item, ph_type: str = "title") -> None:
    """执行归位：写入标题占位符（必要时新建）并删除原文本框。不做统计。"""
    _pt, _top, text, shape = item
    title_ph = find_title_placeholder(slide)
    if title_ph is None:
        title_ph = add_title_placeholder(slide, ph_type)
    tf = title_ph.text_frame
    tf.text = text
    # 不复制原字号：标题字号由版式/母版定义，占位符自动继承
    parent = shape._element.getparent()
    if parent is not None:
        parent.remove(shape._element)


def main() -> int:
    ap = argparse.ArgumentParser(description="Step 1 规范化：标题归位 + 字体继承")
    ap.add_argument("pptx", help="输入 .pptx")
    ap.add_argument("--out", required=True, help="输出 .pptx")
    ap.add_argument("--dry-run", action="store_true", help="只报告不改文件")
    args = ap.parse_args()

    src = Path(args.pptx)
    if not src.exists():
        print("ERROR: 文件不存在 %s" % src, file=sys.stderr)
        return 1

    prs = Presentation(str(src))
    report = {
        "file": str(src),
        "slide_count": len(prs.slides._sldIdLst),
        "placeholder_added": 0,
        "textbox_removed": 0,
        "fonts_cleared": 0,
        "title_moved": [],
        "skipped": [],
        "dry_run": args.dry_run,
    }
    font_counter = [0]

    # 阶段一：制定计划（dry-run 与实跑共用同一套判定，保证数字一致）
    plan = []
    for idx, slide in enumerate(prs.slides, 1):
        title_ph = find_title_placeholder(slide)
        has_text = bool(title_ph is not None and (title_ph.text_frame.text or "").strip())
        try:
            lname = slide.slide_layout.name
        except Exception:
            lname = "?"
        zone = title_zone_for(slide)
        ph_type = title_ph_type_for(slide)
        item = find_title_textbox(slide, has_text, zone)
        plan.append((idx, slide, item, ph_type))
        if item is None:
            if zone is None:
                reason = "版式「%s」无标题槽位，文字保留原位" % lname
            elif has_text:
                reason = "已有标题占位符且含文字"
            else:
                reason = "未找到疑似标题文本框"
            report["skipped"].append({"slide": idx, "reason": reason})
            continue
        pt, _top, text, _shape = item
        report["title_moved"].append({
            "slide": idx, "text": text, "font_pt": pt, "zone": zone,
            "into": ("已有 title 占位符" if title_ph is not None
                     else "新建 %s 占位符" % ph_type),
        })
        if title_ph is None:
            report["placeholder_added"] += 1
        report["textbox_removed"] += 1

    # 阶段二：字体统计（排除即将被删除的文本框，避免虚高）
    doomed = {id(it[2][3]._element) for it in plan if it[2] is not None}
    for _idx, slide, _item, _ptype in plan:
        for shp in iter_shapes(slide.shapes):
            if id(shp._element) in doomed:
                continue
            for tf in iter_text_frames(shp):
                clear_hardcoded_fonts(tf._txBody, font_counter)

    report["fonts_cleared"] = font_counter[0]

    # 阶段三：落盘
    out = Path(args.out)
    if not args.dry_run:
        for _idx, slide, item, ptype in plan:
            if item is not None:
                process_item(slide, item, ptype)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))

    json_path = out.with_suffix("").as_posix() + "_normalize.json"
    Path(json_path).write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print("[normalize] 页数 %d" % report["slide_count"])
    print("[normalize] 标题归位 %d 页 | 新建占位符 %d 个 | 删文本框 %d 个 | 清除硬编码字体 %d 处"
          % (len(report["title_moved"]), report["placeholder_added"],
             report["textbox_removed"], report["fonts_cleared"]))
    if report["skipped"]:
        print("[normalize] 跳过 %d 页（需人工确认）:" % len(report["skipped"]))
        for s in report["skipped"]:
            print("    P%-3d %s" % (s["slide"], s["reason"]))
    print("[normalize] 清单 %s" % json_path)
    if args.dry_run:
        print("[normalize] dry-run，未写出文件")
    else:
        print("[normalize] 产出 %s" % out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
